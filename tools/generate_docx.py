# -*- coding: utf-8 -*-
"""Build a styled DOCX from the SynapseMind Markdown design document."""
import io, os, re
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "docs", "design", "release", "SynapseMind-单题诊断引导-设计方案.md")
DST = os.path.join(ROOT, "docs", "design", "release", "SynapseMind-单题诊断引导-设计方案.docx")
ASSETS = os.path.join(ROOT, "docs", "design", "assets")

# ---------- palette ----------
INK      = RGBColor(0x1F, 0x2A, 0x2E)
MUTED    = RGBColor(0x5D, 0x6B, 0x70)
TEAL     = RGBColor(0x0F, 0x6B, 0x63)
TEAL_HEX = "0F6B63"
AMBER    = RGBColor(0xB0, 0x6A, 0x12)
AMBER_HEX= "B06A12"
CORAL    = RGBColor(0xA8, 0x43, 0x35)
BLUE     = RGBColor(0x2D, 0x5F, 0x8B)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BG = "F7F4EC"
TEAL_BG  = "E8F2EF"
AMBER_BG = "FBF0DC"
CODE_BG  = "F3EFE6"
BORDER   = "D8D2C4"

CN_BODY = "宋体"
CN_HEAD = "微软雅黑"
MONO = "Consolas"

# ---------- low level helpers ----------
def set_run_font(run, name=CN_BODY, size=10.5, color=INK, bold=None, italic=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = color
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:eastAsia"), CN_HEAD if name == CN_HEAD else (MONO if name == MONO else CN_BODY))

def insert_ordered(parent, el, successors):
    """Insert el before the first child whose tag is in successors, else append."""
    for child in list(parent):
        if child.tag in successors:
            parent.insert(list(parent).index(child), el)
            return
    parent.append(el)

def shade_run(run, fill):
    rPr = run._element.get_or_add_rPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    insert_ordered(rPr, shd, {qn(t) for t in (
        "w:fitText","w:vertAlign","w:rtl","w:cs","w:em","w:lang",
        "w:eastAsianLayout","w:specVanish","w:oMath")})

def shade_para(p, fill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    insert_ordered(pPr, shd, {qn(t) for t in (
        "w:tabs","w:suppressAutoHyphens","w:kinsoku","w:wordWrap","w:overflowPunct",
        "w:topLinePunct","w:autoSpaceDE","w:autoSpaceDN","w:bidi","w:adjustRightInd",
        "w:snapToGrid","w:spacing","w:ind","w:contextualSpacing","w:mirrorIndents",
        "w:suppressOverlap","w:jc","w:textDirection","w:textAlignment",
        "w:textboxTightWrap","w:outlineLvl","w:divId","w:cnfStyle","w:rPr",
        "w:sectPr","w:pPrChange")})

def border_para(p, side="bottom", color=TEAL_HEX, sz=8, space=4):
    pPr = p._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        insert_ordered(pPr, pBdr, {qn(t) for t in (
            "w:shd","w:tabs","w:suppressAutoHyphens","w:kinsoku","w:wordWrap",
            "w:overflowPunct","w:topLinePunct","w:autoSpaceDE","w:autoSpaceDN",
            "w:bidi","w:adjustRightInd","w:snapToGrid","w:spacing","w:ind",
            "w:contextualSpacing","w:mirrorIndents","w:suppressOverlap","w:jc",
            "w:textDirection","w:textAlignment","w:textboxTightWrap","w:outlineLvl",
            "w:divId","w:cnfStyle","w:rPr","w:sectPr","w:pPrChange")})
    el = OxmlElement(f"w:{side}")
    el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), str(space)); el.set(qn("w:color"), color)
    pBdr.append(el)

def keep_with_next(p):
    p.paragraph_format.keep_with_next = True

def add_hyperlink(paragraph, text, url, size=10.5):
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), CN_BODY); rFonts.set(qn("w:hAnsi"), CN_BODY); rFonts.set(qn("w:eastAsia"), CN_BODY)
    rPr.append(rFonts)
    color = OxmlElement("w:color"); color.set(qn("w:val"), "2D5F8B"); rPr.append(color)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rPr.append(u)
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(size*2))); rPr.append(sz)
    new_run.append(rPr)
    t = OxmlElement("w:t"); t.set(qn("xml:space"), "preserve"); t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)

TOKEN_RE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`|\[[^\]]+?\]\([^)]+?\))")

def add_inline(p, text, size=10.5, color=INK, base_bold=False):
    text = text.replace("&nbsp;", " ").replace("<br>", "\n").replace("<br/>", "\n").replace("<br />", "\n")
    parts = re.split(TOKEN_RE, text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            r = p.add_run(part[2:-2]); set_run_font(r, CN_BODY, size, color, bold=True)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2 and not part.startswith("**"):
            r = p.add_run(part[1:-1]); set_run_font(r, CN_BODY, size, color, bold=base_bold, italic=True)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            r = p.add_run(part[1:-1]); set_run_font(r, MONO, max(8.0, size-1.2), color, bold=base_bold)
            shade_run(r, "EFEAE0")
        elif part.startswith("[") and "](" in part and part.endswith(")"):
            m = re.match(r"\[([^\]]+?)\]\(([^)]+?)\)", part)
            if m:
                add_hyperlink(p, m.group(1), m.group(2), size)
            else:
                r = p.add_run(part); set_run_font(r, CN_BODY, size, color, bold=base_bold)
        else:
            r = p.add_run(part); set_run_font(r, CN_BODY, size, color, bold=base_bold)

def add_page_field(p, instr, size=9):
    r1 = p.add_run(); set_run_font(r1, CN_BODY, size, MUTED)
    fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin"); r1._r.append(fld1)
    r2 = p.add_run(); set_run_font(r2, CN_BODY, size, MUTED)
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {instr} "
    r2._r.append(it)
    r3 = p.add_run(); set_run_font(r3, CN_BODY, size, MUTED)
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "separate"); r3._r.append(fld2)
    r4 = p.add_run("1"); set_run_font(r4, CN_BODY, size, MUTED)
    r5 = p.add_run(); set_run_font(r5, CN_BODY, size, MUTED)
    fld3 = OxmlElement("w:fldChar"); fld3.set(qn("w:fldCharType"), "end"); r5._r.append(fld3)

def add_toc_field(p):
    r1 = p.add_run(); set_run_font(r1, CN_BODY, 10.5, MUTED)
    fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin"); r1._r.append(fld1)
    r2 = p.add_run(); set_run_font(r2, CN_BODY, 10.5, MUTED)
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
    it.text = ' TOC \\o "1-2" \\h \\z \\u '
    r2._r.append(it)
    r3 = p.add_run(); set_run_font(r3, CN_BODY, 10.5, MUTED)
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "separate"); r3._r.append(fld2)
    r4 = p.add_run("目录将在打开时自动更新；如未更新，请右键选择“更新域”。")
    set_run_font(r4, CN_BODY, 9.5, MUTED, italic=True)
    r5 = p.add_run(); set_run_font(r5, CN_BODY, 10.5, MUTED)
    fld3 = OxmlElement("w:fldChar"); fld3.set(qn("w:fldCharType"), "end"); r5._r.append(fld3)

# ---------- document setup ----------
doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(2.2)
sec.bottom_margin = Cm(2.1)
sec.left_margin = Cm(2.4)
sec.right_margin = Cm(2.4)
sec.header_distance = Cm(1.1)
sec.footer_distance = Cm(1.1)
USABLE = sec.page_width - sec.left_margin - sec.right_margin
USABLE_CM = Emu(USABLE).cm

# Normal
ns = doc.styles["Normal"]
ns.font.name = CN_BODY
ns.font.size = Pt(10.5)
ns.font.color.rgb = INK
ns.paragraph_format.line_spacing = 1.6
ns.paragraph_format.space_after = Pt(6)
ns.element.rPr.rFonts.set(qn("w:eastAsia"), CN_BODY)

# Headings
for name, size, color, before, after, bold in [
    ("Heading 1", 22, TEAL, 0, 14, True),
    ("Heading 2", 14.5, INK, 16, 6, True),
    ("Heading 3", 11.5, TEAL, 10, 4, True),
    ("Heading 4", 10.5, MUTED, 8, 2, True),
]:
    st = doc.styles[name]
    st.font.name = CN_HEAD
    st.font.size = Pt(size)
    st.font.bold = bold
    st.font.color.rgb = color
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True
    st.element.rPr.rFonts.set(qn("w:eastAsia"), CN_HEAD)
doc.styles["Heading 1"].paragraph_format.page_break_before = True
doc.styles["Heading 1"].paragraph_format.space_before = Pt(0)

# A non-outline heading used only for "目录" so it does not list itself in the TOC.
toc_head = doc.styles.add_style("Front TOC Heading", WD_STYLE_TYPE.PARAGRAPH)
toc_head.base_style = doc.styles["Normal"]
toc_head.font.name = CN_HEAD
toc_head.font.size = Pt(22)
toc_head.font.bold = True
toc_head.font.color.rgb = TEAL
toc_head.paragraph_format.space_before = Pt(0)
toc_head.paragraph_format.space_after = Pt(14)
toc_head.paragraph_format.page_break_before = True
toc_head.paragraph_format.keep_with_next = True
toc_head.element.rPr.rFonts.set(qn("w:eastAsia"), CN_HEAD)

front = doc.styles.add_style("Front Section", WD_STYLE_TYPE.PARAGRAPH)
front.base_style = doc.styles["Normal"]
front.font.name = CN_HEAD
front.font.size = Pt(22)
front.font.bold = True
front.font.color.rgb = TEAL
front.paragraph_format.space_before = Pt(0)
front.paragraph_format.space_after = Pt(14)
front.paragraph_format.page_break_before = True
front.paragraph_format.keep_with_next = True
front.element.rPr.rFonts.set(qn("w:eastAsia"), CN_HEAD)

# Header / footer
sec.different_first_page_header_footer = True
hp = sec.header.paragraphs[0]
hp.paragraph_format.tab_stops.add_tab_stop(USABLE, WD_TAB_ALIGNMENT.RIGHT)
r = hp.add_run("SynapseMind 突触智学 · 单题诊断引导"); set_run_font(r, CN_HEAD, 8.5, MUTED)
r = hp.add_run("\tv0.3"); set_run_font(r, CN_HEAD, 8.5, MUTED)
border_para(hp, "bottom", BORDER, 4, 2)

fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run("第 "); set_run_font(r, CN_BODY, 9, MUTED)
add_page_field(fp, "PAGE")
r = fp.add_run(" 页"); set_run_font(r, CN_BODY, 9, MUTED)

# ---------- cover ----------
def para(text="", size=10.5, color=INK, bold=False, italic=False, align=None, before=0, after=0, line=1.0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    if align is not None: p.alignment = align
    if text:
        r = p.add_run(text); set_run_font(r, CN_HEAD if bold else CN_BODY, size, color, bold=bold, italic=italic)
    return p

para("PROJECT DESIGN  ·  2026", 10, TEAL, True, False, WD_ALIGN_PARAGRAPH.LEFT, before=90, after=18)
p = para("SynapseMind", 44, INK, True, False, WD_ALIGN_PARAGRAPH.LEFT, before=0, after=2)
p = para("突触智学 · 单题诊断引导", 30, TEAL, True, False, WD_ALIGN_PARAGRAPH.LEFT, before=0, after=16)
p = para("输入一道题 + 你的答案，找到你卡住的那一步", 15, MUTED, False, False, WD_ALIGN_PARAGRAPH.LEFT, before=0, after=10)
p = para("不是讲这道题，是找到你卡住的那一步。", 14, AMBER, False, True, WD_ALIGN_PARAGRAPH.LEFT, before=0, after=18)
border_para(p, "bottom", TEAL_HEX, 12, 10)

# abstract box
box = doc.add_table(rows=1, cols=1)
box.alignment = WD_TABLE_ALIGNMENT.LEFT
box.autofit = False
box.columns[0].width = Cm(15.4)
cell = box.cell(0, 0)
cell.width = Cm(15.4)
shade_para(cell.paragraphs[0], TEAL_BG)
cp = cell.paragraphs[0]
cp.paragraph_format.space_before = Pt(10)
cp.paragraph_format.space_after = Pt(10)
cp.paragraph_format.left_indent = Cm(0.4)
cp.paragraph_format.right_indent = Cm(0.4)
r = cp.add_run("把“我不会”从一句情绪，变成一个可命名、可验证、可训练的卡点。")
set_run_font(r, CN_HEAD, 12, TEAL, bold=True)
cp2 = cell.add_paragraph()
cp2.paragraph_format.space_before = Pt(0)
cp2.paragraph_format.space_after = Pt(10)
cp2.paragraph_format.left_indent = Cm(0.4)
cp2.paragraph_format.right_indent = Cm(0.4)
r = cp2.add_run("学生发来一道题 + 自己的答案；系统用 2–6 轮追问定位具体卡点，再给 3 分钟微干预与一道迁移题。")
set_run_font(r, CN_BODY, 10.5, INK)

para("", 10, INK, before=0, after=0)
para("高中数学 · 函数与导数    |    100 题 · 30 题深标    |    3 人 · 8 周    |    3 个容器 · 2 种语言", 10, MUTED, False, False, WD_ALIGN_PARAGRAPH.LEFT, before=18, after=0)
para("项目设计方案  v0.3    ·    2026 年 10 月", 10, MUTED, False, False, WD_ALIGN_PARAGRAPH.LEFT, before=6, after=0)
para("", 10, INK, before=0, after=0)
para("本文为设计方案，不是已上线产品。所有指标均为设计目标或估算；缺口见第 9.8 节。", 9, MUTED, False, True, WD_ALIGN_PARAGRAPH.LEFT, before=30, after=0)

doc.add_page_break()

# ---------- markdown parsing ----------
raw = io.open(SRC, encoding="utf-8").read().splitlines()
start = next(i for i, l in enumerate(raw) if l.strip() == "## 目录")
lines = raw[start:]

def is_table_sep(s):
    cells = [c.strip() for c in s.strip().strip("|").split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", c) for c in cells)

def split_table_row(s):
    s = s.strip()
    if s.startswith("|"): s = s[1:]
    if s.endswith("|") and not s.endswith("\\|"): s = s[:-1]
    parts = re.split(r"(?<!\\)\|", s)
    return [c.replace("\\|", "|").strip() for c in parts]

def is_bold_only(s):
    s = s.strip()
    return s.startswith("**") and s.endswith("**") and s.count("**") == 2 and len(s) > 4

def add_body(text, indent=True, size=10.5, color=INK):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.6
    p.paragraph_format.space_after = Pt(6)
    if indent:
        p.paragraph_format.first_line_indent = Pt(21)
    add_inline(p, text, size, color)
    return p

def add_bullet(text, level=0):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.55 + 0.55*level)
    p.paragraph_format.first_line_indent = Cm(-0.45)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.45
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run("•  " if level == 0 else "–  ")
    set_run_font(r, CN_BODY, 10.5, TEAL, bold=True)
    add_inline(p, text, 10.5, INK)

def add_number(text, n):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.7)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.45
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(f"{n}.  ")
    set_run_font(r, CN_BODY, 10.5, TEAL, bold=True)
    add_inline(p, text, 10.5, INK)

def add_quote(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.35)
    p.paragraph_format.right_indent = Cm(0.2)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.45
    shade_para(p, TEAL_BG)
    border_para(p, "left", TEAL_HEX, 18, 8)
    add_inline(p, text, 9.8, MUTED)

def add_code_block(code_lines, lang=""):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.25)
    p.paragraph_format.right_indent = Cm(0.2)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.0
    shade_para(p, CODE_BG)
    border_para(p, "left", "C9C2B4", 10, 6)
    for i, line in enumerate(code_lines):
        if i > 0:
            p.add_run().add_break()
        r = p.add_run(line if line else " ")
        set_run_font(r, MONO, 8.0, RGBColor(0x33, 0x3A, 0x3D))

def add_mermaid_sequence(code_lines):
    # parse participants
    names = {}
    steps = []
    for line in code_lines:
        s = line.strip()
        m = re.match(r"participant\s+(\S+)\s+as\s+(.+)", s)
        if m:
            names[m.group(1)] = m.group(2).strip()
            continue
        m = re.match(r"(\S+?)\s*->>\s*(\S+?)\s*:\s*(.+)", s)
        if m:
            a, b, msg = m.group(1), m.group(2), m.group(3)
            a = names.get(a, a); b = names.get(b, b)
            if a == b:
                steps.append(f"{a}：{msg}")
            else:
                steps.append(f"{a} → {b}：{msg}")
            continue
        m = re.match(r"Note over\s+([^:]+):\s*(.+)", s)
        if m:
            steps.append(f"注（{m.group(1).strip()}）：{m.group(2).strip()}")
            continue
        if s.startswith("alt "):
            steps.append(f"【条件】{s[4:].strip()}")
        elif s.startswith("else "):
            steps.append(f"【否则】{s[5:].strip()}")
        elif s.startswith("loop "):
            steps.append(f"【循环】{s[5:].strip()}")
        elif s == "end":
            steps.append("【结束】")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("时序步骤"); set_run_font(r, CN_HEAD, 10.5, TEAL, bold=True)
    r = p.add_run("（Mermaid 源码见 Markdown 版）"); set_run_font(r, CN_BODY, 8.5, MUTED)
    n = 0
    for st in steps:
        if st.startswith("【"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.55)
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(st); set_run_font(r, CN_BODY, 9.2, AMBER, bold=True)
        elif st.startswith("注"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.7)
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(st); set_run_font(r, CN_BODY, 9.0, MUTED, italic=True)
        else:
            n += 1
            add_number(st, n)

def add_mermaid_note(kind):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.35)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    shade_para(p, "F7F4EC")
    note = {
        "graph": "【架构/分层图】Mermaid 源码见 Markdown 版；总览见图 5-1，Agent 编排见图 6-1。",
        "state": "【状态机图】Mermaid 源码见 Markdown 版；角色与编排总览见图 6-1，交接规则见 6.3.3。",
    }.get(kind, "【图】Mermaid 源码见 Markdown 版。")
    r = p.add_run(note); set_run_font(r, CN_BODY, 9.0, MUTED, italic=True)

def add_table(rows):
    if not rows:
        return
    header = rows[0]
    body = [r for r in rows[1:] if not is_table_sep("|".join(r))]
    ncol = max(len(r) for r in rows)
    header = header + [""] * (ncol - len(header))
    body = [r + [""] * (ncol - len(r)) for r in body]

    table = doc.add_table(rows=1, cols=ncol)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.allow_autofit = False

    # widths by average length
    weights = []
    for c in range(ncol):
        vals = [len(r[c]) for r in rows if c < len(r)]
        avg = sum(vals) / max(1, len(vals))
        weights.append(max(1.2, min(avg, 26)))
    total = sum(weights)
    widths = [max(1.1, USABLE_CM * w / total) for w in weights]
    # normalize
    scale = USABLE_CM / sum(widths)
    widths = [w * scale for w in widths]

    # table width 100%
    tblPr = table._tbl.tblPr
    tblW = OxmlElement("w:tblW"); tblW.set(qn("w:type"), "pct"); tblW.set(qn("w:w"), "5000")
    insert_ordered(tblPr, tblW, {qn(t) for t in (
        "w:jc","w:tblCellSpacing","w:tblInd","w:tblBorders","w:shd","w:tblLayout",
        "w:tblCellMar","w:tblLook","w:tblCaption","w:tblDescription")})
    # borders
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:space"), "0"); el.set(qn("w:color"), BORDER)
        borders.append(el)
    insert_ordered(tblPr, borders, {qn(t) for t in (
        "w:shd","w:tblLayout","w:tblCellMar","w:tblLook","w:tblCaption","w:tblDescription")})

    # header
    hdr = table.rows[0]
    trPr = hdr._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true")
    insert_ordered(trPr, th, {qn(t) for t in ("w:tblCellSpacing","w:jc","w:hidden","w:ins","w:del","w:trPrChange")})
    for i, txt in enumerate(header):
        cell = hdr.cells[i]
        cell.width = Cm(widths[i])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        shade_para(cell.paragraphs[0], TEAL_HEX)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(1.5)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.line_spacing = 1.1
        r = p.add_run(txt.replace("**", ""))
        set_run_font(r, CN_HEAD, 8.2, WHITE, bold=True)

    # body
    for ri, row in enumerate(body):
        tr = table.add_row()
        trPr = tr._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit"); cant.set(qn("w:val"), "true")
        insert_ordered(trPr, cant, {qn(t) for t in (
            "w:trHeight","w:tblHeader","w:tblCellSpacing","w:jc","w:hidden",
            "w:ins","w:del","w:trPrChange")})
        for ci, txt in enumerate(row):
            cell = tr.cells[ci]
            cell.width = Cm(widths[ci])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if ri % 2 == 1:
                shade_para(cell.paragraphs[0], "F7F9F8")
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1.2)
            p.paragraph_format.space_after = Pt(1.2)
            p.paragraph_format.line_spacing = 1.12
            add_inline(p, txt, 8.2, INK)

    # cell margins
    for row in table.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcMar = OxmlElement("w:tcMar")
            for side, val in (("top", 40), ("left", 70), ("bottom", 40), ("right", 70)):
                el = OxmlElement(f"w:{side}"); el.set(qn("w:w"), str(val)); el.set(qn("w:type"), "dxa"); tcMar.append(el)
            insert_ordered(tcPr, tcMar, {qn(t) for t in (
                "w:textDirection","w:tcFitText","w:vAlign","w:hideMark",
                "w:headers","w:cellIns","w:cellDel","w:cellMerge","w:tcPrChange")})

    # spacer
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(2)

# ---------- main loop ----------
i = 0
skip_toc = False
num_counter = 0
while i < len(lines):
    line = lines[i]
    s = line.strip()

    if not s:
        i += 1
        continue

    if skip_toc:
        if s.startswith("## "):
            skip_toc = False
        else:
            i += 1
            continue

    # A numbered list only continues while consecutive lines are numbered.
    if not re.match(r"^\d+\.\s+", s) or is_bold_only(s):
        num_counter = 0

    # headings
    if s.startswith("#### "):
        p = doc.add_paragraph(style="Heading 3")
        add_inline(p, s[5:], 11.5, TEAL, base_bold=True)
        i += 1; continue
    if s.startswith("### "):
        p = doc.add_paragraph(style="Heading 2")
        add_inline(p, s[4:], 14.5, INK, base_bold=True)
        i += 1; continue
    if s.startswith("## "):
        title = s[3:].strip()
        if title == "目录":
            p = doc.add_paragraph(style="Front TOC Heading")
            add_inline(p, title, 22, TEAL, base_bold=True)
            border_para(p, "bottom", TEAL_HEX, 10, 6)
            tp = doc.add_paragraph()
            tp.paragraph_format.space_before = Pt(8)
            add_toc_field(tp)
            skip_toc = True
            i += 1; continue
        if title in ("阅读提示", "一张图看懂"):
            p = doc.add_paragraph(style="Front Section")
            add_inline(p, title, 22, TEAL, base_bold=True)
            border_para(p, "bottom", TEAL_HEX, 10, 6)
            i += 1; continue
        p = doc.add_paragraph(style="Heading 1")
        add_inline(p, title, 22, TEAL, base_bold=True)
        border_para(p, "bottom", TEAL_HEX, 10, 6)
        i += 1; continue
    if s.startswith("# "):
        # only if a stray H1 remains
        p = doc.add_paragraph(style="Heading 1")
        add_inline(p, s[2:], 22, TEAL, base_bold=True)
        border_para(p, "bottom", TEAL_HEX, 10, 6)
        i += 1; continue

    # horizontal rule
    if s == "---":
        i += 1; continue

    # code fence
    if s.startswith("```"):
        lang = s[3:].strip().lower()
        code = []
        i += 1
        while i < len(lines) and not lines[i].strip().startswith("```"):
            code.append(lines[i])
            i += 1
        i += 1
        if lang == "mermaid":
            first = next((x.strip() for x in code if x.strip()), "")
            if first.startswith("sequenceDiagram"):
                add_mermaid_sequence(code)
            elif first.startswith("stateDiagram"):
                add_mermaid_note("state")
            else:
                add_mermaid_note("graph")
        else:
            add_code_block(code, lang)
        continue

    # image
    m = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", s)
    if m:
        alt, path = m.group(1), m.group(2)
        full = os.path.join(os.path.dirname(SRC), path.replace("/", os.sep))
        if os.path.exists(full):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run()
            run.add_picture(full, width=USABLE)
            # caption: next italic line
            cap = None
            if i + 1 < len(lines):
                nxt = lines[i+1].strip()
                if nxt.startswith("*") and nxt.endswith("*") and len(nxt) > 2:
                    cap = nxt.strip("*").strip()
                    i += 1
            if cap:
                cp = doc.add_paragraph()
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cp.paragraph_format.space_before = Pt(0)
                cp.paragraph_format.space_after = Pt(8)
                r = cp.add_run(cap); set_run_font(r, CN_BODY, 8.8, MUTED, italic=True)
        else:
            add_body(f"[缺失图片：{path}]", indent=False, color=CORAL)
        i += 1; continue

    # table
    if s.startswith("|") and i + 1 < len(lines) and lines[i+1].strip().startswith("|"):
        rows = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            rows.append(split_table_row(lines[i]))
            i += 1
        add_table(rows)
        continue

    # blockquote
    if s.startswith(">"):
        q = []
        while i < len(lines) and lines[i].strip().startswith(">"):
            q.append(lines[i].strip()[1:].strip())
            i += 1
        add_quote(" ".join(q))
        continue

    # lists
    if re.match(r"^[-*]\s+", s):
        level = 0
        if line.startswith("  - ") or line.startswith("    - "):
            level = 1
        add_bullet(re.sub(r"^[-*]\s+", "", s), level)
        num_counter = 0
        i += 1; continue
    if re.match(r"^\d+\.\s+", s):
        num_counter += 1
        add_number(re.sub(r"^\d+\.\s+", "", s), num_counter)
        i += 1; continue

    # bold-only callout label
    if is_bold_only(s) and len(s) <= 28:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        keep_with_next(p)
        r = p.add_run(s.strip("*").strip())
        set_run_font(r, CN_HEAD, 11, TEAL, bold=True)
        i += 1; continue

    # paragraph: accumulate consecutive text lines
    buf = [s]
    i += 1
    while i < len(lines):
        nxt = lines[i].strip()
        if (not nxt or nxt.startswith("#") or nxt.startswith("|") or nxt.startswith("```")
            or nxt.startswith(">") or re.match(r"^[-*]\s+", nxt) or re.match(r"^\d+\.\s+", nxt)
            or nxt.startswith("![") or nxt == "---"):
            break
        buf.append(nxt)
        i += 1
    # smart join: space only between ASCII alnum boundaries
    text = buf[0]
    for nxt in buf[1:]:
        if text and nxt and re.match(r"[A-Za-z0-9)]", text[-1]) and re.match(r"[A-Za-z0-9(]", nxt[0]):
            text += " " + nxt
        else:
            text += nxt
    add_body(text)

# ---------- save ----------
doc.core_properties.title = "SynapseMind 突触智学 · 单题诊断引导"
doc.core_properties.author = "SynapseMind"
doc.core_properties.subject = "输入一道题与自己的答案，通过苏格拉底式追问定位具体卡点"
doc.core_properties.keywords = "单题诊断, 苏格拉底追问, 认知诊断, 多Agent, RAG, Spring Cloud, 教育产品"
doc.core_properties.comments = "v0.3 单题诊断引导方案；技术选型见 ADR-0001；指标为设计目标或估算。"

doc.save(DST)
print("saved", DST)
print("size", os.path.getsize(DST))